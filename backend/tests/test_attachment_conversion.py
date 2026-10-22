"""What an uploaded file's text is, and how the file tree stores every readable type of file."""

from __future__ import annotations

import asyncio
import io
import zipfile
from pathlib import Path

import docx
import openpyxl
import pytest

from tianzhou_agent_platform.aina.document.builtin import READ_DOCUMENT_TOOL_ID, invoke_document_tool
from tianzhou_agent_platform.aina.document.service import DocumentService
from tianzhou_agent_platform.conversations import conversion
from tianzhou_agent_platform.conversations.conversion import convert
from tianzhou_agent_platform.core.errors import PlatformError
from tianzhou_agent_platform.store.errors import StorageValidationError
from tianzhou_agent_platform.store.nas.filesystem import NasStore

ACTOR = {"user_id": "anonymous", "tenant_id": "default"}


def pdf(*pages: str) -> bytes:
    """A PDF whose pages carry ``pages`` as text (an empty string is a page without a text layer)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", "", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {len(objects)} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    return out + f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def word(*, heading: str, body: str, table: list[list[str]]) -> bytes:
    document = docx.Document()
    document.add_heading(heading, level=2)
    document.add_paragraph(body)
    grid = document.add_table(rows=len(table), cols=len(table[0]))
    for row, values in zip(grid.rows, table, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def workbook(sheets: dict[str, list[list[object]]]) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def presentation(*slides: list[str]) -> bytes:
    """A PPTX whose slides hold ``slides`` as paragraphs; the slide files are named in reverse of the deck's order."""
    main = "http://schemas.openxmlformats.org/presentationml/2006/main"
    drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
    relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ids = "".join(f'<p:sldId id="{256 + index}" r:id="rId{index}"/>' for index in range(1, len(slides) + 1))
    targets = "".join(
        f'<Relationship Id="rId{index}" Target="slides/slide{len(slides) + 1 - index}.xml"/>'
        for index in range(1, len(slides) + 1)
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{main}" xmlns:r="{relationships}"><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>',
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{targets}</Relationships>',
        )
        for index, paragraphs in enumerate(slides, start=1):
            # A paragraph's text may be split across runs.
            body = "".join(
                f"<a:p>{''.join(f'<a:r><a:t>{part}</a:t></a:r>' for part in text.split('|'))}</a:p>" for text in paragraphs
            )
            archive.writestr(
                f"ppt/slides/slide{len(slides) + 1 - index}.xml",
                f'<p:sld xmlns:p="{main}" xmlns:a="{drawing}"><p:cSld><p:spTree>{body}</p:spTree></p:cSld></p:sld>',
            )
    return buffer.getvalue()


def refused(name: str, payload: bytes) -> PlatformError:
    with pytest.raises(PlatformError) as caught:
        convert(name, payload)
    return caught.value


def test_a_text_file_is_its_own_text() -> None:
    notes = convert("notes.md", b"# Notes\n\nHello")
    log = convert("run.log", "started".encode("utf-16"))  # with a byte order mark
    table = convert("data.csv", b"a,b\n1,2\n")

    assert (notes.text, notes.content_type, notes.binary) == ("# Notes\n\nHello", "text/markdown", False)
    assert log.text == "started" and log.content_type == "text/plain"
    assert (table.text, table.content_type) == ("a,b\n1,2\n", "text/csv")


def test_files_unibot_cannot_read_are_refused() -> None:
    image = refused("photo.png", b"\x89PNG\r\n")
    binary = refused("dump.txt", b"abc\x00def")
    latin1 = refused("old.txt", "café".encode("latin-1"))
    fake_pdf = refused("report.pdf", b"just text")
    fake_docx = refused("spec.docx", b"PK\x03\x04 not an archive")
    no_extension = refused("README", b"text")

    assert (image.status_code, image.code) == (415, "UNSUPPORTED_FILE_TYPE")
    assert no_extension.status_code == 415
    assert all(error.status_code == 422 for error in (binary, latin1, fake_pdf, fake_docx))
    assert "not UTF-8 text" in latin1.message


def test_a_pdf_becomes_one_section_per_page() -> None:
    converted = convert("report.pdf", pdf("Revenue grew", "", "Costs fell"))

    assert converted.content_type == "application/pdf" and converted.binary
    assert converted.text == "## Page 1\n\nRevenue grew\n\n## Page 2\n\n## Page 3\n\nCosts fell\n"


def test_a_scanned_pdf_without_a_text_layer_is_refused() -> None:
    error = refused("scan.pdf", pdf("", ""))

    assert error.status_code == 422 and "no text layer" in error.message


def test_a_docx_keeps_its_headings_and_tables() -> None:
    converted = convert(
        "spec.docx", word(heading="Scope", body="Only the API.", table=[["Name", "Owner"], ["API", "a|b"]])
    )

    assert converted.text == ("## Scope\n\nOnly the API.\n\n| Name | Owner |\n| --- | --- |\n| API | a\\|b |\n")


def test_an_xlsx_becomes_one_csv_section_per_sheet() -> None:
    converted = convert(
        "budget.xlsx",
        workbook({"Q1": [["Item", "Cost"], ["Rent", 1200], [None, None], ["Food", 300.5], [None]], "Empty": []}),
    )

    assert converted.text == (
        "## Sheet: Q1\n\n```csv\nItem,Cost\nRent,1200\n\nFood,300.5\n```\n\n## Sheet: Empty\n\n(empty)\n"
    )


def test_a_pptx_becomes_one_section_per_slide_in_the_deck_order() -> None:
    converted = convert("deck.pptx", presentation(["Apples", "A quick |look"], [], ["Thank you", " "]))
    fake = refused("deck.pptx", b"PK\x03\x04 not an archive")

    assert converted.binary
    assert converted.content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    # A slide of pictures alone keeps its heading.
    assert converted.text == "## Slide 1\n\nApples\nA quick look\n\n## Slide 2\n\n## Slide 3\n\nThank you\n"
    assert fake.status_code == 422


def test_an_office_file_that_unpacks_to_too_much_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    payload = word(heading="Big", body="x" * 5000, table=[["a"]])
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        unpacked = sum(item.file_size for item in archive.infolist())
    monkeypatch.setattr(conversion, "MAX_UNPACKED_BYTES", unpacked - 1)

    assert "unpacks to more than" in refused("big.docx", payload).message


def test_text_larger_than_a_document_may_hold_is_refused() -> None:
    error = refused("huge.txt", b"x" * (1024 * 1024 + 1))

    assert (error.status_code, error.code) == (413, "QUOTA_EXCEEDED")


def test_an_upload_is_stored_in_uploads_under_its_own_name(tmp_path: Path) -> None:
    service = DocumentService(NasStore(tmp_path))

    async def scenario() -> tuple[list[str], list[str], bytes, str]:
        names = [
            await service.add_upload("report.pdf", b"%PDF-1", "## Page 1\n\nOne", **ACTOR),
            await service.add_upload("report.pdf", b"%PDF-2", "## Page 1\n\nTwo", **ACTOR),
            await service.add_upload("notes.md", b"# Notes", None, **ACTOR),
            await service.add_upload("a:b?<con>.txt", b"text", None, **ACTOR),
            await service.add_upload("con.txt", b"text", None, **ACTOR),
        ]
        listed = [item.name for item in await service.list_documents(**ACTOR)]
        second = await service.get_document("Uploads/report (2).pdf", **ACTOR)
        return names, listed, await service.read_file("Uploads/report (2).pdf", **ACTOR), second.content

    names, listed, original, text = asyncio.run(scenario())

    assert names == [
        "Uploads/report.pdf",
        "Uploads/report (2).pdf",
        "Uploads/notes.md",
        "Uploads/a_b__con_.txt",
        "Uploads/con_.txt",
    ]
    assert listed == sorted(names, key=str.casefold)  # the hidden extracted text is not listed
    assert (original, text) == (b"%PDF-2", "## Page 1\n\nTwo")


def test_renaming_or_deleting_a_pdf_carries_its_extracted_text(tmp_path: Path) -> None:
    service = DocumentService(NasStore(tmp_path))
    uploads = tmp_path / "documents" / "t-default" / "u-anonymous" / "Uploads"

    async def scenario() -> tuple[bytes, str]:
        await service.add_upload("report.pdf", b"%PDF-1", "## Page 1\n", **ACTOR)
        await service.rename_document("Uploads/report.pdf", "Uploads/q3.pdf", **ACTOR)
        moved = (
            await service.read_file("Uploads/q3.pdf", **ACTOR),
            (await service.get_document("Uploads/q3.pdf", **ACTOR)).content,
        )
        with pytest.raises(StorageValidationError, match="keeps its extension"):
            await service.rename_document("Uploads/q3.pdf", "Uploads/q3.txt", **ACTOR)
        await service.delete_document("Uploads/q3.pdf", **ACTOR)
        return moved

    assert asyncio.run(scenario()) == (b"%PDF-1", "## Page 1\n")
    assert list(uploads.iterdir()) == []


def test_what_each_kind_of_file_allows(tmp_path: Path) -> None:
    service = DocumentService(NasStore(tmp_path))

    async def refusal(call: object) -> str:
        with pytest.raises(StorageValidationError) as caught:
            await call  # type: ignore[misc]
        return str(caught.value)

    async def scenario() -> dict[str, object]:
        await service.add_upload("report.pdf", b"%PDF-1", "## Page 1\n\nRevenue\n", **ACTOR)
        await service.create_document("tool.py", "# a comment, not a heading\nprint(1)\n", **ACTOR)
        tool = await service.get_document("tool.py", **ACTOR)
        saved = await service.update_sections("tool.py", "print(2)\n", _revision(tool.content), **ACTOR)
        return {
            "pdf_outline": [
                item.heading for item in (await service.get_outline("Uploads/report.pdf", **ACTOR)).headings
            ],
            "pdf_edit": await refusal(service.update_document("Uploads/report.pdf", "x", **ACTOR)),
            "pdf_create": await refusal(service.create_document("new.pdf", "x", **ACTOR)),
            "py_outline": (await service.get_outline("tool.py", **ACTOR)).headings,
            "py_read_section": await refusal(service.get_section("tool.py", "a", 1, **ACTOR)),
            "py_section": await refusal(service.update_section("tool.py", "a", 1, "# a", "r", **ACTOR)),
            "py_saved": (saved.updated_sections, (await service.get_document("tool.py", **ACTOR)).content),
            "png_name": await refusal(service.create_document("photo.png", "x", **ACTOR)),
        }

    result = asyncio.run(scenario())

    assert result["pdf_outline"] == ["Page 1"]
    assert "read-only" in result["pdf_edit"] and "only uploaded" in result["pdf_create"]
    assert result["py_outline"] == [] and "has no sections" in result["py_read_section"]
    assert "has no sections to edit" in result["py_section"]
    assert result["py_saved"] == ([], "print(2)\n")
    assert "Only Markdown, text, code, PDF, DOCX, XLSX and PPTX" in result["png_name"]


def _revision(content: str) -> str:
    from tianzhou_agent_platform.aina.document.service import _document_revision

    return _document_revision(content)


def test_document_read_pages_through_lines(tmp_path: Path) -> None:
    service = DocumentService(NasStore(tmp_path))

    async def read(**arguments: object) -> dict[str, object]:
        result, _widgets = await invoke_document_tool(
            service, READ_DOCUMENT_TOOL_ID, {"name": "long.md", **arguments}, **ACTOR
        )
        return result["document"]

    async def scenario() -> list[dict[str, object]]:
        await service.create_document("long.md", "one\ntwo\nthree\nfour\n", **ACTOR)
        return [await read(), await read(offset=2, limit=2), await read(offset=4, limit=10), await read(offset=9)]

    whole, middle, tail, beyond = asyncio.run(scenario())

    assert whole["content"] == "one\ntwo\nthree\nfour\n" and "lines" not in whole
    assert (middle["content"], middle["lines"]) == ("two\nthree\n", {"start": 2, "end": 3, "total": 4})
    assert (tail["content"], tail["lines"]) == ("four\n", {"start": 4, "end": 4, "total": 4})
    assert (beyond["content"], beyond["lines"]) == ("", {"start": 9, "end": 8, "total": 4})
    with pytest.raises(PlatformError):
        asyncio.run(read(offset=0))
